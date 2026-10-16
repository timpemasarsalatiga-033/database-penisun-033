from flask import Flask,request,redirect,url_for,session,send_file,flash,render_template_string,jsonify
import sqlite3,os,zipfile
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from functools import wraps

app=Flask(__name__)
app.secret_key=os.environ.get("SECRET_KEY","penisun033-change-this")
BASE=os.path.dirname(os.path.abspath(__file__))
DB=os.path.join(BASE,"database.db")
XLSX=os.path.join(BASE,"database_taspen.xlsx")
BACKUP=os.path.join(BASE,"backups")
os.makedirs(BACKUP,exist_ok=True)

TAS={"Nama":"nama","NIK/NIP":"nik","Alamat":"alamat","No HP":"no_hp","Peserta Taspen":"peserta","Pengurusan":"pengurusan","No Rekening":"rekening","Keterangan":"keterangan","No Tiket":"no_tiket","Tanda Terima":"tanda_terima"}
TDS={"Nama":"nama","No HP":"no_hp","NOTAS":"notas","Alamat":"alamat"}
SPTB={"Status":"status","Nama Cabang":"cabang","Nama Kelompok Bayar":"kelompok_bayar","Nama Juru Bayar":"juru_bayar","NOTAS":"notas","No Rekening":"rekening","Nama Peserta":"peserta","Nama Penerima":"penerima"}

STYLE='''<style>
:root{--b:#075fa9;--d:#063b73;--g:#d4a83a;--bg:#f4f7fb;--ink:#18334b}*{box-sizing:border-box}body{margin:0;font-family:Segoe UI,Arial,sans-serif;background:var(--bg);color:var(--ink)}a{text-decoration:none;color:var(--b)}button,.btn{border:0;background:var(--b);color:#fff;border-radius:9px;padding:10px 15px;font-weight:700;cursor:pointer}.secondary{background:#e8eef4;color:var(--ink)}.side{position:fixed;inset:0 auto 0 0;width:235px;background:linear-gradient(180deg,var(--d),var(--b));color:#fff;padding:25px 17px}.side h2{margin:0 8px 30px}.side a{display:block;color:#fff;padding:11px 12px;border-radius:9px;margin:4px 0}.side a:hover{background:#ffffff18}.maker{position:absolute;bottom:18px;padding:0 10px;font-size:11px;opacity:.8}.main{margin-left:235px;padding:25px;min-height:100vh}.hero{background:linear-gradient(120deg,var(--d),#0a75c9);color:#fff;border-radius:20px;padding:35px;min-height:240px;position:relative;overflow:hidden}.hero h1{font-size:34px}.hero span{color:#f1d27c}.hero .art{position:absolute;right:8%;bottom:-25px;font-size:110px;opacity:.2}.cards{display:grid;grid-template-columns:repeat(3,1fr);gap:18px;margin:20px 0}.card,.panel{background:#fff;border-radius:15px;padding:20px;box-shadow:0 8px 30px #15385c0d}.card b{font-size:30px;color:var(--b);display:block}.head{display:flex;justify-content:space-between;align-items:center;margin:15px 0}.filters{display:flex;gap:8px;margin-bottom:15px}.filters input,.filters select,.form input,.form textarea,.form select{border:1px solid #d8e2eb;border-radius:8px;padding:10px;font:inherit;background:#fff}.filters input{flex:1}.table{overflow:auto}.table table{width:100%;min-width:1050px;border-collapse:collapse}.table th,.table td{padding:10px;border-bottom:1px solid #edf1f5;text-align:left;font-size:12px;vertical-align:top}.table th{background:#f6f9fc}.actions{white-space:nowrap}.actions a{margin-right:8px}.danger{color:#c33;background:none;padding:0}.form{display:grid;grid-template-columns:repeat(2,1fr);gap:16px}.form label{display:flex;flex-direction:column;gap:6px;font-weight:700;font-size:13px}.form textarea{min-height:100px}.full{grid-column:1/-1}.login{min-height:100vh;display:grid;place-items:center;background:linear-gradient(135deg,#062f5c,#0c74c4)}.loginbox{background:#fff;padding:45px;border-radius:22px;width:min(420px,92%);box-shadow:0 25px 70px #001a3d66}.loginbox h1 span{color:var(--b)}.loginbox form{display:grid;gap:9px}.alert{padding:12px;border-radius:9px;background:#e7f7ed;margin-bottom:12px}.printbar{padding:12px;text-align:right;background:#fff;position:sticky;top:0}.paper{max-width:190mm;margin:auto}.receipt{background:#fff;border:1px solid #222;padding:20px;margin:20px}.row{display:grid;grid-template-columns:42% 58%;padding:7px;border-bottom:1px solid #ddd}@media(max-width:850px){.side{position:static;width:100%;}.maker{display:none}.main{margin:0;padding:14px}.cards{grid-template-columns:1fr}.form{grid-template-columns:1fr}.full{grid-column:auto}.filters{flex-wrap:wrap}.filters input{flex-basis:100%}.hero .art{display:none}}
</style>'''

def db():
    c=sqlite3.connect(DB);c.row_factory=sqlite3.Row;return c

def init():
    c=db();c.executescript("""CREATE TABLE IF NOT EXISTS taspen(id INTEGER PRIMARY KEY AUTOINCREMENT,no INTEGER,nama TEXT,nik TEXT,alamat TEXT,no_hp TEXT,peserta TEXT,pengurusan TEXT,rekening TEXT,keterangan TEXT,no_tiket TEXT,tanda_terima TEXT,created_at TEXT,updated_at TEXT);
    CREATE TABLE IF NOT EXISTS tds(id INTEGER PRIMARY KEY AUTOINCREMENT,nama TEXT,no_hp TEXT,notas TEXT,alamat TEXT,created_at TEXT,updated_at TEXT);
    CREATE TABLE IF NOT EXISTS sptb(id INTEGER PRIMARY KEY AUTOINCREMENT,status TEXT,cabang TEXT,kelompok_bayar TEXT,juru_bayar TEXT,notas TEXT,rekening TEXT,peserta TEXT,penerima TEXT,created_at TEXT,updated_at TEXT);""");c.commit();c.close()

def export_xlsx():
    c=db();w=Workbook();s=w.active;s.title="Tanda Terima Taspen"
    for name,m,table in [("Tanda Terima Taspen",TAS,"taspen"),("Apply TDS",TDS,"tds"),("SPTB",SPTB,"sptb")]:
        if name!="Tanda Terima Taspen": s=w.create_sheet(name)
        cols=list(m.values());s.append(list(m.keys()))
        for r in c.execute("SELECT "+",".join(cols)+" FROM "+table+" ORDER BY id"):s.append([r[x] for x in cols])
        for cell in s[1]:cell.font=Font(bold=True,color="FFFFFF");cell.fill=PatternFill("solid",fgColor="075FA9");cell.alignment=Alignment(horizontal="center")
        s.freeze_panes="A2"
    w.save(XLSX);c.close()

def backup():
    export_xlsx();p=os.path.join(BACKUP,"database_penisun_033_"+datetime.now().strftime("%Y%m%d_%H%M%S")+".zip")
    with zipfile.ZipFile(p,"w",zipfile.ZIP_DEFLATED) as z:z.write(DB,"database.db");z.write(XLSX,"database_taspen.xlsx")
    return p

def auth(f):
    @wraps(f)
    def w(*a,**k):
        if not session.get("user"):return redirect(url_for("login"))
        return f(*a,**k)
    return w

def layout(title,body):
    nav='<aside class="side"><h2>DATABASE<br>PENISUN 033</h2><a href="/">Dashboard</a><a href="/taspen">Tanda Terima Taspen</a><a href="/tds">Apply TDS</a><a href="/sptb">SPTB</a><a href="/export">Export Excel</a><a href="/backup">Backup Database</a><a href="/logout">Keluar</a><div class="maker">Aplikasi ini dibuat oleh OkkyIndra</div></aside>'
    return "<!doctype html><html lang=id><meta name=viewport content='width=device-width,initial-scale=1'><title>"+title+"</title>"+STYLE+nav+"<main class=main>"+body+"</main></html>"

@app.route("/login",methods=["GET","POST"])
def login():
    if request.method=="POST":
        if request.form.get("username")=="BJT" and request.form.get("password")=="033":session["user"]="BJT";return redirect("/")
        flash("Username atau password salah.")
    msg="".join("<div class='alert'>"+x+"</div>" for x in session.pop("_flashes",[]))
    return "<!doctype html><html lang=id><meta name=viewport content='width=device-width,initial-scale=1'>"+STYLE+"<body class=login><section class=loginbox><h1>DATABASE <span>PENISUN 033</span></h1><p>Digital Data Management System</p>"+msg+"<form method=post><label>Username</label><input name=username required><label>Password</label><input name=password type=password required><button>MASUK</button></form><p><small>Aplikasi ini dibuat oleh OkkyIndra</small></p></section></body></html>"

@app.route("/logout")
def logout():session.clear();return redirect("/login")

@app.route("/")
@auth
def home():
    c=db();a=[c.execute("SELECT COUNT(*) FROM "+t).fetchone()[0] for t in ("taspen","tds","sptb")];c.close()
    body=f'<section class=hero><h1>Kinerja & Database<br><span>Pensiun Bank Jateng</span></h1><p>DATABASE PENISUN 033 — pusat pengelolaan data digital.</p><div class=art>⌨</div></section><div class=cards><div class=card><b>{a[0]}</b>Tanda Terima Taspen</div><div class=card><b>{a[1]}</b>Apply TDS</div><div class=card><b>{a[2]}</b>SPTB</div></div><div class=panel><h3>Akses Cepat</h3><a class=btn href="/taspen">Kelola Taspen</a> <a class=btn href="/tds">Kelola TDS</a> <a class=btn href="/sptb">Kelola SPTB</a></div>'
    return layout("Dashboard",body)

def list_view(kind):
    maps={"taspen":TAS,"tds":TDS,"sptb":SPTB};m=maps[kind];q=request.args.get("q","");status=request.args.get("status","");c=db();sql="SELECT * FROM "+kind+" WHERE 1=1";args=[]
    fields={"taspen":["nama","nik","no_hp","rekening","no_tiket"],"tds":["nama","no_hp","notas","alamat"],"sptb":["notas","rekening","peserta","penerima","cabang"]}[kind]
    if q:sql+=" AND ("+" OR ".join(x+" LIKE ?" for x in fields)+")";args += ["%"+q+"%"]*len(fields)
    if kind=="sptb" and status:sql+=" AND status=?";args.append(status)
    rows=c.execute(sql+" ORDER BY id DESC",args).fetchall();c.close()
    heads=list(m.keys());html=f'<div class=head><h1>{dict(taspen="Tanda Terima Taspen",tds="Apply TDS",sptb="SPTB")[kind]}</h1><a class=btn href="/{kind}/new">+ Tambah Data</a></div><form class=filters><input name=q value="{q}" placeholder="Cari data...">'
    if kind=="sptb":html+=f'<select name=status><option value="">Semua Status</option><option {"selected" if status=="Belum Lapor" else ""}>Belum Lapor</option><option {"selected" if status=="Sudah Lapor" else ""}>Sudah Lapor</option></select>'
    html+='<button>Cari</button></form><div class="panel table"><table><tr>'+''.join("<th>"+x+"</th>" for x in heads)+"<th>Aksi</th></tr>"
    for r in rows:
        html+="<tr>"+''.join("<td>"+str(r[m[k]] or "-")+"</td>" for k in heads)+f'<td class=actions><a href="/{kind}/{r["id"]}/edit">Edit</a><a href="/print/{kind}/{r["id"]}">Print</a><form style="display:inline" method=post action="/{kind}/{r["id"]}/delete" onsubmit="return confirm(\'Hapus data?\')"><button class=danger>Hapus</button></form></td></tr>'
    html+="</table></div>";return layout(dict(taspen="Taspen",tds="TDS",sptb="SPTB")[kind],html)

for k in ("taspen","tds","sptb"):
    app.add_url_rule("/"+k,k,lambda k=k:list_view(k))

@app.route("/<kind>/new",methods=["GET","POST"])
@auth
def new(kind):
    return form_view(kind,None)

@app.route("/<kind>/<int:id>/edit",methods=["GET","POST"])
@auth
def edit(kind,id):
    c=db();row=c.execute("SELECT * FROM "+kind+" WHERE id=?",(id,)).fetchone();c.close();return form_view(kind,row)

def form_view(kind,row):
    m={"taspen":TAS,"tds":TDS,"sptb":SPTB}[kind]
    if request.method=="POST":
        data={v:request.form.get(v,"").strip() for v in m.values()};now=datetime.now().isoformat(timespec="seconds");c=db()
        if kind=="taspen" and not row:data["no"]=c.execute("SELECT COALESCE(MAX(no),0)+1 FROM taspen").fetchone()[0]
        cols=list(data);vals=[data[x] for x in cols]
        if row:c.execute("UPDATE "+kind+" SET "+",".join(x+"=?" for x in cols)+",updated_at=? WHERE id=?",vals+[now,row["id"]])
        else:c.execute("INSERT INTO "+kind+" ("+",".join(cols)+",created_at,updated_at) VALUES ("+(",".join("?" for _ in cols))+",?,?)",vals+[now,now])
        c.commit();c.close();export_xlsx();return redirect("/"+kind)
    fields=list(m.items());html=f'<h1>{"Edit" if row else "Tambah"} {kind.upper()}</h1><form class=form method=post>'
    if kind=="taspen":html+=f'<label>Nomor otomatis<input value="{row["no"] if row else "Otomatis"}" disabled></label>'
    for label,key in fields:
        if kind=="taspen" and label=="No":continue
        val=row[key] if row else ""
        if label=="Alamat" or label=="Keterangan":html+=f'<label class=full>{label}<textarea name="{key}">{val or ""}</textarea></label>'
        elif kind=="sptb" and label=="Status":html+=f'<label>{label}<select name="{key}"><option {"selected" if val=="Belum Lapor" else ""}>Belum Lapor</option><option {"selected" if val=="Sudah Lapor" else ""}>Sudah Lapor</option></select></label>'
        else:html+=f'<label>{label}<input name="{key}" value="{val or ""}"></label>'
    html+='<div class=full><button>Simpan</button> <a class="btn secondary" href="/'+kind+'">Batal</a></div></form>';return layout("Form",html)

@app.post("/<kind>/<int:id>/delete")
@auth
def delete(kind,id):
    c=db();c.execute("DELETE FROM "+kind+" WHERE id=?",(id,));c.commit();c.close();export_xlsx();return redirect("/"+kind)

@app.route("/print/<kind>/<int:id>")
@auth
def print_page(kind,id):
    c=db();r=c.execute("SELECT * FROM "+kind+" WHERE id=?",(id,)).fetchone();c.close();m={"taspen":TAS,"tds":TDS,"sptb":SPTB}[kind]
    rows="".join(f'<div class=row><b>{k}</b><span>{r[v] or "-"}</span></div>' for k,v in m.items() if not (kind=="taspen" and k=="Tanda Terima"))
    return "<!doctype html><html><meta name=viewport content='width=device-width,initial-scale=1'>"+STYLE+f'<div class=printbar><button onclick="window.print()">Cetak / PDF</button></div><div class=paper><div class=receipt><h2 style="text-align:center">BANK JATENG CABANG SALATIGA</h2><h2 style="text-align:center">TANDA TERIMA DATA</h2>{rows}<div style="margin-top:60px;display:flex;justify-content:space-between"><span>Petugas</span><span>Penerima</span></div></div></div></html>'

@app.route("/export")
@auth
def export():export_xlsx();return send_file(XLSX,as_attachment=True,download_name="database_taspen.xlsx")

@app.route("/backup")
@auth
def backup_route():
    return send_file(backup(),as_attachment=True,download_name=os.path.basename(backup()))

@app.route("/api/stats")
@auth
def stats():
    c=db();x={t:c.execute("SELECT COUNT(*) FROM "+t).fetchone()[0] for t in ("taspen","tds","sptb")};c.close();return jsonify(x)

init()
if __name__=="__main__":app.run(host="0.0.0.0",port=int(os.environ.get("PORT",5033)))
